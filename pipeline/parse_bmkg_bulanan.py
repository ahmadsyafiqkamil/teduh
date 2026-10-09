"""Parse BMKG Data Online monthly XLSX exports into one tidy CSV (DATA_SOURCES.md §E).

The XLSX cannot be fetched unattended: login uses reCAPTCHA and downloading is gated on a
per-semester satisfaction survey. Drop the files from dataonline.bmkg.go.id into
data/raw/bmkg/bulanan/ by hand, then run this.

Run: uv run --no-project --python 3.12 --with openpyxl python pipeline/parse_bmkg_bulanan.py
Output: data/raw/bmkg/stations_monthly.csv (one row per station-month)
"""

import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "raw" / "bmkg" / "bulanan"
OUT = ROOT / "data" / "raw" / "bmkg" / "stations_monthly.csv"

# BMKG sentinels, documented in the KETERANGAN block of every sheet.
MISSING = {"-", "", "8888", "9999", "8888.0", "9999.0"}

# BMKG column header -> our name. wind_dir/wind_dir_car are compass text, kept as-is.
COLUMNS = {
    "rainfall": "rainfall_mm",
    "wind_speed": "wind_max_ms",
    "wind_avg": "wind_avg_ms",
    "hum_avg": "rh_avg_pct",
    "sunshine": "sunshine_h",
    "temp_max": "tmax_c",
    "temp_min": "tmin_c",
    "temp_avg": "tavg_c",
    "wind_dir": "wind_dir",
    "wind_dir_car": "wind_dir_car",
}
NUMERIC = ("rainfall_mm", "wind_max_ms", "wind_avg_ms", "rh_avg_pct", "sunshine_h",
           "tmax_c", "tmin_c", "tavg_c")
META = ("station_id", "station_name", "lat", "lon", "elev_m", "period")
FIELDS = list(META) + list(COLUMNS.values())

HEADER_CELL = "Periode"
META_LABELS = {"ID WMO": "station_id", "NAMA STASIUN": "station_name",
               "LINTANG": "lat", "BUJUR": "lon", "ELEVASI": "elev_m"}


def clean(value) -> str | None:
    """Return None for BMKG's missing-data sentinels, else a trimmed string."""
    text = "" if value is None else str(value).strip()
    return None if text in MISSING else text


def read_meta(rows: list[tuple]) -> dict:
    meta = {}
    for row in rows:
        cells = [clean(c) for c in row]
        label = next((c for c in cells if c and c.rstrip(":") in META_LABELS), None)
        if not label:
            continue
        values = [c for c in cells if c and c != label and c.strip(": ")]
        if values:
            meta[META_LABELS[label.rstrip(":")]] = values[-1].lstrip(": ").removesuffix(" Meter")
    return meta


def read_sheet(rows: list[tuple]) -> tuple[dict, list[dict]]:
    meta = read_meta(rows)
    header_idx = next(i for i, r in enumerate(rows) if clean(r[0]) == HEADER_CELL)
    header = [clean(c) for c in rows[header_idx]]

    records = []
    for row in rows[header_idx + 1:]:
        period = clean(row[0])
        if not period or not period[:4].isdigit():
            break  # KETERANGAN legend starts here
        rec = dict(meta, period=period)
        for col, value in zip(header[1:], row[1:]):
            if col in COLUMNS:
                rec[COLUMNS[col]] = clean(value)
        records.append(rec)
    return meta, records


def dedupe(rows: list[dict]) -> list[dict]:
    """Keep one row per station-month, which is what this file promises its consumers.

    Exports overlap on purpose: the Jun-Sep 2025 request was made before the wider
    2024-11..2026-02 one, and both are kept on disk as the raw record. BMKG has so far
    repeated identical figures for the shared months, but a silent revision would quietly
    bias any cross-check against GHCN, so a disagreement is reported rather than hidden.
    """
    kept: dict[tuple[str, str], dict] = {}
    conflicts = []
    for row in rows:
        key = (row.get("station_id", "?"), row.get("period", "?"))
        first = kept.setdefault(key, row)
        if first is row:
            continue
        differing = {k: (first.get(k), row.get(k)) for k in NUMERIC if first.get(k) != row.get(k)}
        if differing:
            conflicts.append(f"{key[0]} {key[1]}: {differing}")
    if conflicts:
        print(f"PERINGATAN: {len(conflicts)} bulan dengan nilai berbeda antar ekspor, "
              f"dipakai yang pertama:")
        for line in conflicts:
            print(f"  {line}")
    dropped = len(rows) - len(kept)
    if dropped:
        print(f"dedupe    : {dropped} baris ganda dibuang (ekspor tumpang tindih)")
    return list(kept.values())


def main() -> None:
    files = sorted(SRC.glob("*.xlsx"))
    if not files:
        raise SystemExit(f"no xlsx in {SRC.relative_to(ROOT)} — export from dataonline.bmkg.go.id first")

    all_rows, empty = [], []
    for path in files:
        wb = openpyxl.load_workbook(path, data_only=True)
        for name in wb.sheetnames:
            rows = list(wb[name].iter_rows(values_only=True))
            meta, records = read_sheet(rows)
            filled = [r for r in records if any(r.get(k) for k in NUMERIC)]
            label = f"{meta.get('station_id','?')} {meta.get('station_name','?')[:28]}"
            print(f"{path.name}  {label:34s} {len(filled)}/{len(records)} bulan berisi data")
            if not filled:
                empty.append(label)
            all_rows += filled

    all_rows = dedupe(all_rows)
    all_rows.sort(key=lambda r: (r.get("station_id", ""), r.get("period", "")))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(all_rows)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(all_rows)} rows)")
    if empty:
        print(f"PERINGATAN: stasiun tanpa data sama sekali: {', '.join(empty)}")
    check(all_rows)


def check(rows: list[dict]) -> None:
    """One row per station-month, sentinels dropped, temps plausible, tmin <= tavg <= tmax."""
    assert rows, "no rows parsed"
    keys = [(r.get("station_id"), r.get("period")) for r in rows]
    assert len(keys) == len(set(keys)), "station-month ganda lolos ke output; dedupe gagal"
    for r in rows:
        for k in NUMERIC:
            assert r.get(k) not in MISSING or r.get(k) is None, f"sentinel leaked into {k}"
        t = {k: float(r[k]) for k in ("tmin_c", "tavg_c", "tmax_c") if r.get(k)}
        for k, v in t.items():
            assert 15.0 <= v <= 45.0, f"implausible {k}={v} at {r['station_id']} {r['period']}"
        if len(t) == 3:
            assert t["tmin_c"] <= t["tavg_c"] <= t["tmax_c"], f"temp order wrong at {r['period']}"
        if r.get("rh_avg_pct"):
            assert 0.0 <= float(r["rh_avg_pct"]) <= 100.0, f"bad RH at {r['period']}"
    print("check ok")


if __name__ == "__main__":
    main()
